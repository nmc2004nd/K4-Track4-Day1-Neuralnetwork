"""results_table.py — Lưu JSON và tạo bảng tổng hợp thí nghiệm.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    exp_id = result.get("cfg", {}).get("exp_id")
    if not exp_id:
        raise ValueError("result['cfg']['exp_id'] không được để trống")

    output_dir = Path(results_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{exp_id}.json"
    payload = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"],
    }
    with output_path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    return str(output_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    results = []
    for path in Path(results_dir).glob("*.json"):
        with path.open("r", encoding="utf-8") as handle:
            result = json.load(handle)
        if not all(key in result for key in ("cfg", "history", "summary")):
            raise ValueError(f"File không đúng định dạng kết quả: {path}")
        if not result["cfg"].get("exp_id"):
            raise ValueError(f"File thiếu cfg.exp_id: {path}")
        results.append(result)
    return sorted(results, key=lambda result: result["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    summary = result["summary"]
    hidden = cfg.get("hidden")
    if isinstance(hidden, (list, tuple)):
        hidden = "-".join(str(width) for width in hidden)

    row = {
        "exp_id": cfg.get("exp_id"),
        "group": cfg.get("group"),
        "description": cfg.get("description"),
        "loss": cfg.get("loss"),
        "optimizer": cfg.get("optimizer"),
        "lr": cfg.get("lr"),
        "weight_decay": cfg.get("weight_decay"),
        "batch": cfg.get("batch"),
        "epochs": cfg.get("epochs"),
        "hidden": hidden,
        "dropout": cfg.get("dropout"),
        "clip_norm": cfg.get("clip_norm") if cfg.get("clip_norm") is not None else "none",
        "precision": cfg.get("precision"),
        "init": cfg.get("init"),
        "seed": cfg.get("seed"),
        "step0_loss": summary.get("step0_loss"),
        "best_val_loss": summary.get("best_val_loss"),
        "best_epoch": summary.get("best_epoch"),
        "final_train_loss": summary.get("final_train_loss"),
        "final_val_loss": summary.get("final_val_loss"),
        "val_acc": summary.get("val_acc"),
        "val_macro_f1": summary.get("val_macro_f1"),
        "time_per_epoch_s": summary.get("time_per_epoch_s"),
        "peak_mem_MB": summary.get("peak_mem_MB"),
        "diverged": summary.get("diverged"),
        "eval_acc": None,
        "eval_macro_f1": None,
        "figure_file": f"figures/{cfg.get('exp_id')}.png",
        "notes": notes,
    }
    if eval_scores is not None:
        row["eval_acc"] = eval_scores.get("accuracy", eval_scores.get("eval_acc"))
        row["eval_macro_f1"] = eval_scores.get(
            "macro_f1", eval_scores.get("eval_macro_f1")
        )
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    import openpyxl

    formula_columns = {
        "step0_gap_vs_lnC",
        "gap_val_minus_train",
        "delta_val_f1_vs_base",
        "beyond_noise",
    }
    workbook = openpyxl.load_workbook(template_path, data_only=False)
    worksheet = workbook["Experiments"]
    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }

    writable_headers = [name for name in headers if name not in formula_columns]
    capacity = worksheet.max_row - 1
    if len(rows) > capacity:
        raise ValueError(f"Template chỉ có {capacity} dòng, nhưng nhận được {len(rows)} kết quả")

    # Xoá dữ liệu mẫu/cũ ở các cột nhập tay; giữ nguyên công thức và định dạng.
    for row_index in range(2, worksheet.max_row + 1):
        for header in writable_headers:
            worksheet.cell(row=row_index, column=headers[header]).value = None

    for row_index, row in enumerate(rows, start=2):
        for key, value in row.items():
            if key in headers and key not in formula_columns:
                worksheet.cell(row=row_index, column=headers[key]).value = value

    output_path = Path(out_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
